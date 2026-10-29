from flask import Flask, request, jsonify, send_file, render_template_string, abort, session, redirect, url_for
import os
import threading
import time
import openpyxl
from urllib.parse import quote
import io
import json
import drive_helper
import tempfile
import urllib.parse

app = Flask(__name__, static_folder='static')
app.secret_key = 'super_secret_archive_key'

ROOT_DRIVE_FOLDER_NAME = "DATA SHEET INDEX"
OVERVIEW_TITLE = "R2 PROCESSING AREA"
JUNCTION_FOLDER_NAME = "JUNCTION BOXES"

all_data = []
data_lock = threading.Lock()

# ============ دوال المساعدة ============
def build_hyperlink_map(ws):
    hmap = {}
    try:
        for hl in ws._hyperlinks:
            if hl.ref and hl.target:
                hmap[hl.ref] = hl.target
    except Exception:
        pass
    return hmap

def get_cell_hyperlink_target(cell, hmap):
    if cell.hyperlink and cell.hyperlink.target:
        return cell.hyperlink.target
    target = hmap.get(cell.coordinate)
    if target:
        return target
    return None

# ============ المسح في الخلفية ============
def scan_drive_recursive(service, folder_id, current_path=""):
    files = drive_helper.list_files_in_folder(service, folder_id)
    for f in files:
        if f['mimeType'] == 'application/vnd.google-apps.folder':
            scan_drive_recursive(service, f['id'], current_path + f['name'] + "/")
        elif f['name'].endswith(('.xlsx', '.xlsm')) and not f['name'].startswith('~$'):
            try:
                file_bytes = drive_helper.download_file(service, f['id'])
                wb = openpyxl.load_workbook(file_bytes, data_only=True)
                ws = wb.active
                hmap = build_hyperlink_map(ws)
                
                for row in ws.iter_rows():
                    for cell in row:
                        target = get_cell_hyperlink_target(cell, hmap)
                        if cell.value and target:
                            with data_lock:
                                all_data.append({
                                    "tag": str(cell.value).strip(),
                                    "path": current_path + f['name'],
                                    "file_id": f['id'],
                                    "folder": current_path.split('/')[0] if current_path else "Root"
                                })
            except Exception as e:
                print(f"Error reading {f['name']}: {e}")

def background_scan():
    while True:
        service = drive_helper.get_drive_service()
        if service:
            print("Scanning Google Drive... Please wait.")
            root_id = drive_helper.find_folder_id(service, ROOT_DRIVE_FOLDER_NAME)
            if root_id:
                with data_lock:
                    all_data.clear()
                scan_drive_recursive(service, root_id)
                print(f"Scan complete. Found {len(all_data)} tags.")
            else:
                print(f"Root folder '{ROOT_DRIVE_FOLDER_NAME}' not found in Google Drive.")
        else:
            print("Waiting for authorization...")
        time.sleep(3600)

# ============ API Routes ============
@app.route("/api/stats")
def api_stats():
    with data_lock:
        data = list(all_data)
    folders = {}
    for d in data:
        folders[d["folder"]] = folders.get(d["folder"], 0) + 1
    return jsonify({
        "total": len(data),
        "folder_count": len(folders),
        "excel_files": len(set(d["path"] for d in data)),
        "folders": folders
    })

@app.route("/api/search")
def api_search():
    q = request.args.get("q", "").lower()
    with data_lock:
        data = list(all_data)
    results = []
    for i, d in enumerate(data):
        if q in d["tag"].lower():
            results.append({"id": i, "tag": d["tag"], "sheet_no": "", "file_id": d["file_id"]})
    return jsonify(results)

@app.route("/file/<file_id>")
def serve_drive_file(file_id):
    service = drive_helper.get_drive_service()
    if not service: abort(401)
    try:
        file_meta = service.files().get(fileId=file_id, fields="name, mimeType").execute()
        file_bytes = drive_helper.download_file(service, file_id)
        temp_path = os.path.join(tempfile.gettempdir(), file_meta['name'])
        with open(temp_path, 'wb') as f:
            f.write(file_bytes.getvalue())
        return send_file(temp_path, mimetype=file_meta['mimeType'])
    except Exception as e:
        return str(e), 500

@app.route("/api/xlsx-table")
def api_xlsx_table():
    file_id = request.args.get("file_id")
    if not file_id: abort(404)
    
    service = drive_helper.get_drive_service()
    if not service: return jsonify({"error": "Not authorized"}), 401
    
    try:
        file_meta = service.files().get(fileId=file_id, fields="name, parents").execute()
        excel_parent_id = file_meta['parents'][0] if 'parents' in file_meta else None
        file_bytes = drive_helper.download_file(service, file_id)
        wb = openpyxl.load_workbook(file_bytes)
        ws = wb.active
        hmap = build_hyperlink_map(ws)
    except Exception as e:
        print(f"❌ ERROR downloading/reading file {file_id}: {e}")
        return jsonify({"error": str(e)}), 500

    raw_rows = list(ws.iter_rows())
    clean_grid = []
    
    for row in raw_rows:
        clean_row = []
        for cell in row:
            href = None
            target = get_cell_hyperlink_target(cell, hmap)
            if target:
                decoded_target = urllib.parse.unquote(target.replace("\\", "/"))
                if '#' in decoded_target:
                    decoded_target = decoded_target.split('#')[0]
                href = f"/find-and-open?target={quote(decoded_target)}&parent_id={excel_parent_id}"
                    
            val = "" if cell.value is None else str(cell.value).strip()
            if val or href:
                style = {}
                try:
                    if cell.fill.patternType == "solid":
                        fill_color = cell.fill.start_color
                        if fill_color.rgb and fill_color.rgb != "00000000": style["bg"] = f"#{fill_color.rgb[2:]}"
                except: pass
                clean_row.append({"value": val, "href": href, "style": style})
        if clean_row: clean_grid.append(clean_row)
        
    return jsonify({"rows": clean_grid})

@app.route("/find-and-open")
def find_and_open():
    target = request.args.get("target")
    parent_id = request.args.get("parent_id")
    if not target: abort(400)
    
    service = drive_helper.get_drive_service()
    if not service: abort(401)
    
    path_parts = [p for p in target.split('/') if p and p != '..']
    filename = path_parts[-1]
    linked_file_id = None
    
    if parent_id and len(path_parts) > 1:
        linked_file_id = drive_helper.find_file_by_path(service, path_parts, parent_id)
    if not linked_file_id and parent_id:
        linked_file_id = drive_helper.find_file_in_tree(service, filename, parent_id)
    if not linked_file_id:
        linked_file_id = drive_helper.find_file_by_name(service, filename)
        
    if linked_file_id:
        return redirect(f"/file/{linked_file_id}")
    else:
        return f"<h1>File Not Found</h1><p>The file '{filename}' was not found.</p>", 404

@app.route("/api/junction-data")
def api_junction_data():
    service = drive_helper.get_drive_service()
    if not service: return jsonify({"error": "Unauthorized"}), 401
    root_id = drive_helper.find_folder_id(service, ROOT_DRIVE_FOLDER_NAME)
    if not root_id: return jsonify({"categories": []})
    folder_id = drive_helper.find_folder_id(service, JUNCTION_FOLDER_NAME, parent_id=root_id)
    if not folder_id: return jsonify({"categories": []})
    files = drive_helper.list_files_in_folder(service, folder_id)
    items = []
    for f in files:
            if f['name'].lower().endswith(('.xlsx', '.xlsm')):
                items.append({"tag": f['name'].split('.')[0], "href": f['id']})
        
    categories = [{"label": "Junction Boxes Files", "items": items}] if items else []
    return jsonify({"categories": categories})
@app.route("/api/overview-data")
def api_overview_data():
    service = drive_helper.get_drive_service()
    if not service: return jsonify({"error": "Unauthorized"}), 401
    root_id = drive_helper.find_folder_id(service, ROOT_DRIVE_FOLDER_NAME)
    if not root_id: return jsonify({"error": "Root not found"}), 404
    pi_id = drive_helper.find_folder_id(service, "P&I", parent_id=root_id)
    if not pi_id: return jsonify({"error": "P&I folder not found"}), 404
    files = drive_helper.list_files_in_folder(service, pi_id)
    pi_file_id = next((f['id'] for f in files if f['name'].lower() == 'p&i.xlsx'), None)
    if not pi_file_id: return jsonify({"error": "P&I.xlsx not found"}), 404
    
    try:
        file_bytes = drive_helper.download_file(service, pi_file_id)
        wb = openpyxl.load_workbook(file_bytes, data_only=True)
        ws = wb.active
        hmap = build_hyperlink_map(ws)
        categories = []
        for row in ws.iter_rows():
            current_label = None
            current_items = []
            def flush():
                if current_label: categories.append({"label": current_label, "items": current_items})
            for cell in row:
                target = get_cell_hyperlink_target(cell, hmap)
                if target:
                    if current_label:
                        filename = os.path.basename(target.replace("\\", "/"))
                        linked_id = drive_helper.find_file_by_name(service, filename)
                        href = f"/file/{linked_id}" if linked_id else None
                        current_items.append({"tag": str(cell.value).strip() if cell.value else "?", "href": href})
                elif cell.value not in (None, ""):
                    flush(); current_label = str(cell.value).strip(); current_items = []
                else:
                    flush(); current_label = None; current_items = []
            flush()
        return jsonify({"categories": categories})
    except Exception as e:
        return jsonify({"error": str(e)}), 500
@app.route("/api/drive-browse")
def api_drive_browse():
    service = drive_helper.get_drive_service()
    if not service: return jsonify({"error": "Unauthorized"}), 401

    folder_id = request.args.get("folder_id")
    
    try:
        if not folder_id:
            folder_id = drive_helper.find_folder_id(service, ROOT_DRIVE_FOLDER_NAME)
            if not folder_id: return jsonify({"error": "Root folder not found"}), 404

        files = drive_helper.list_files_in_folder(service, folder_id)
        folders_out = []
        files_out = []

        for f in files:
            if f['mimeType'] == 'application/vnd.google-apps.folder':
                folders_out.append({"id": f['id'], "name": f['name']})
            else:
                ext = os.path.splitext(f['name'])[1].lower()
                if ext in ('.xlsx', '.xlsm', '.pdf', '.jpg', '.jpeg', '.png', '.drawio'):
                    files_out.append({"id": f['id'], "name": f['name'], "ext": ext})

        folders_out.sort(key=lambda x: x['name'])
        files_out.sort(key=lambda x: x['name'])

        return jsonify({"folders": folders_out, "files": files_out})
    except Exception as e:
        print(f"❌ Browse Error: {e}")
        return jsonify({"error": str(e)}), 500
# ============ القوالب (HTML) ============
PAGE_HTML = """
<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>Control and Systems Archive</title>
<script src="/static/echarts.min.js"></script>
<script src="/static/echarts-gl.min.js"></script>
<style>
  html, body { height:100%; margin:0; }
  body { background:#1a1a1a; color:#eee; font-family: Tahoma, Arial, sans-serif; display:flex; height:100vh; }
  #sidebar { width:300px; background:#232323; padding:16px; box-sizing:border-box; display:flex; flex-direction:column; }
  #sidebar h1 { font-size:16px; margin:0 0 16px; }
  input[type=text] { width:100%; padding:8px; border-radius:6px; border:none; background:#2b2b2b; color:#fff; box-sizing:border-box; }
  button { cursor:pointer; border:none; border-radius:6px; background:#1f538d; color:#fff; padding:8px 12px; margin-top:8px; }
  button:hover { background:#2a6bb0; }
  .tabs { display:flex; gap:6px; margin-top:12px; }
  .tab { flex:1; background:#333; text-align:center; }
  .tab.active { background:#1f538d; }
  #results { flex:1; overflow-y:auto; margin-top:10px; background:#2b2b2b; border-radius:6px; }
  .row { padding:8px 10px; border-bottom:1px solid #333; cursor:pointer; font-size:13px; }
  .row:hover { background:#1f538d; }
  #main { flex:1; display:flex; flex-direction:column; padding:16px; box-sizing:border-box; overflow:hidden; }
  .cards { display:flex; gap:10px; margin-bottom:16px; flex-shrink:0; }
  .card { flex:1; border-radius:10px; padding:20px; text-align:center; color:#fff; font-weight:bold; font-size:20px; }
  .card b { display:block; font-size:42px; margin-top:8px; }
  .card:nth-child(1) { background:#00bcd4; }
  .card:nth-child(2) { background:#2ecc71; }
  .card:nth-child(3) { background:#ffa726; }
  .chart-container { flex-grow:1; background:#2b2b2b; border-radius:8px; overflow:hidden; position:relative; }
  #folder-chart { width:100%; height:100%; }
  #viewer { display: none; flex:1; background:#000; border-radius:8px; flex-direction:column; overflow:hidden; }
  #viewer iframe { width:100%; height:100%; border:none; }
  body.focus-mode .chart-container, body.focus-mode .cards { display: none !important; }
  body.focus-mode #viewer { display: flex; }
  #exit-focus-btn { position:absolute; top:20px; right:20px; display:none; background:#d9534f; z-index:100; }
  body.focus-mode #exit-focus-btn { display: block; }
  table.xlsx { border-collapse:collapse; font-size:18px; margin:auto; width:100%; }
  table.xlsx td { border:1px solid #444; padding:10px; white-space:nowrap; background:#111; color:#eee; }
  table.xlsx a { color:#6cf; text-decoration:underline; cursor:pointer; }
</style>
</head>
<body>
<div id="sidebar">
  <h1>🏭 Archive (Google Drive)</h1>
  <div class="tabs">
    <button class="tab active" id="tab-search" onclick="switchTab('search')">🔍 Search</button>
    <button class="tab" id="tab-browse" onclick="switchTab('browse')">🗂 Browse</button>
  </div>
  <button style="width:100%;" onclick="location.href='/overview'">🗺️ R2 General Map</button>
  <button style="width:100%;" onclick="location.href='/junction-boxes'">🔌 Junction Boxes</button>
  <div id="search-box">
    <input type="text" id="q" placeholder="Search Tag...">
    <button onclick="doSearch()">🔍 Search</button>
  </div>
  <div id="results"></div>
</div>
<div id="main">
  <button id="exit-focus-btn" onclick="exitFocusMode()">⬅ Exit Fullscreen</button>
  <div class="cards">
    <div class="card">Total Tags<b id="c-total">0</b></div>
    <div class="card">Total Folders<b id="c-folders">0</b></div>
    <div class="card">Excel Files<b id="c-excel">0</b></div>
  </div>
  <div class="chart-container"><div id="folder-chart"></div></div>
  <div id="viewer">Select a file to open.</div>
</div>

<script>
let mode = 'search';
let currentFolderPath = [];

function switchTab(which) {
  mode = which;
  document.getElementById('tab-search').classList.toggle('active', which === 'search');
  document.getElementById('tab-browse').classList.toggle('active', which === 'browse');
  document.getElementById('search-box').style.display = which === 'search' ? 'block' : 'none';
  document.getElementById('results').innerHTML = '';
  if (which === 'search') { doSearch(); }
  else { currentFolderPath = []; loadDriveFolder(''); }
}

async function loadDriveFolder(folderId) {
  const box = document.getElementById('results');
  box.innerHTML = 'Loading... Please wait...';
  
  try {
    const r = await fetch(`/api/drive-browse?folder_id=${folderId}`);
    if (!r.ok) {
        throw new Error(`Server responded with status ${r.status}`);
    }
    const data = await r.json();
    
    if (data.error) { 
        box.innerHTML = `<span style="color:#d9534f;">Error: ${data.error}</span>`; 
        return; 
    }
    
    box.innerHTML = '';
    
    // زر الرجوع
    if (currentFolderPath.length > 0) {
      const upBtn = document.createElement('div');
      upBtn.className = 'row';
      upBtn.innerHTML = '⬅ Back';
      upBtn.onclick = () => {
        currentFolderPath.pop();
        const parentId = currentFolderPath.length > 0 ? currentFolderPath[currentFolderPath.length - 1].id : '';
        loadDriveFolder(parentId);
      };
      box.appendChild(upBtn);
    }
    
    // عرض المجلدات
    data.folders.forEach(f => {
      const div = document.createElement('div');
      div.className = 'row';
      div.innerHTML = `📁 ${f.name}`;
      div.onclick = () => {
        currentFolderPath.push({ id: f.id, name: f.name });
        loadDriveFolder(f.id);
      };
      box.appendChild(div);
    });
    
    // عرض الملفات
    data.files.forEach(f => {
      const div = document.createElement('div');
      div.className = 'row';
      const icon = (f.ext === '.xlsx' || f.ext === '.xlsm') ? '📊' : '📄';
      div.innerHTML = `${icon} ${f.name}`;
      if (f.ext === '.xlsx' || f.ext === '.xlsm') {
        div.onclick = () => openXlsxTable(f.id);
      } else {
        div.onclick = () => window.open(`/file/${f.id}`, '_blank');
      }
      box.appendChild(div);
    });
    
    // إذا كان المجلد فارغاً
    if (data.folders.length === 0 && data.files.length === 0) {
        box.innerHTML += '<div class="row" style="color:#888;">Empty folder</div>';
    }

  } catch (e) {
    box.innerHTML = `<span style="color:#d9534f;">Failed to load: ${e.message}</span>`;
  }
}

async function loadStats() {
  const r = await fetch('/api/stats');
  const s = await r.json();
  document.getElementById('c-total').textContent = s.total;
  document.getElementById('c-folders').textContent = s.folder_count;
  document.getElementById('c-excel').textContent = s.excel_files;
  const chartDom = document.getElementById('folder-chart');
  const myChart = echarts.init(chartDom);
  const folders = Object.keys(s.folders).map(name => [name, s.folders[name], s.folders[name]]);
  folders.sort((a, b) => b[1] - a[1]);
  myChart.setOption({
    tooltip: {}, visualMap: { max: Math.max(...Object.values(s.folders)), inRange: { color: ['#00bcd4', '#ffa726'] } },
    xAxis3D: { type: 'category', data: folders.map(f => f[0]), axisLabel: { rotate: 45, color: '#fff', fontSize: 10 } },
    yAxis3D: { type: 'value' }, zAxis3D: { type: 'value' },
    grid3D: { boxWidth: 200, boxDepth: 80, viewControl: { autoRotate: true } },
    series: [{ type: 'bar3D', data: folders, shading: 'lambert', barSize: 10 }]
  });
  window.addEventListener('resize', () => myChart.resize());
}

function addRow(text, onClick) {
  const div = document.createElement('div'); div.className = 'row'; div.textContent = text; div.onclick = onClick;
  document.getElementById('results').appendChild(div);
}

async function doSearch() {
  const q = document.getElementById('q').value.trim();
  const box = document.getElementById('results');
  if (!q) { box.innerHTML = ''; return; }
  const r = await fetch(`/api/search?q=${encodeURIComponent(q)}`);
  const files = await r.json();
  box.innerHTML = '';
  files.forEach(f => addRow(f.tag, () => openXlsxTable(f.file_id)));
}

function enterFocusMode() { document.body.classList.add('focus-mode'); }
function exitFocusMode() { document.body.classList.remove('focus-mode'); }

async function openXlsxTable(fileId) {
  enterFocusMode();
  const viewer = document.getElementById('viewer');
  viewer.innerHTML = 'Loading file from Google Drive...';
  const r = await fetch(`/api/xlsx-table?file_id=${fileId}`);
  const data = await r.json();
  if (data.error) { viewer.textContent = 'Error: ' + data.error; return; }
  viewer.innerHTML = '';
  const searchContainer = document.createElement('div');
  searchContainer.style.cssText = 'background: #222; padding: 10px;';
  const searchInput = document.createElement('input');
  searchInput.type = 'text'; searchInput.placeholder = '🔍 Search inside this file...';
  searchInput.style.cssText = 'width: 100%; padding: 10px; background: #333; color: #fff; border: none;';
  searchContainer.appendChild(searchInput);
  viewer.appendChild(searchContainer);
  const table = document.createElement('table');
  table.className = 'xlsx';
  data.rows.forEach(row => {
    const tr = document.createElement('tr');
    row.forEach(cell => {
      const td = document.createElement('td');
      if (cell.style && cell.style.bg) td.style.backgroundColor = cell.style.bg;
      if (cell.href) {
        const a = document.createElement('a');
        a.textContent = cell.value;
        a.href = cell.href;
        a.target = '_blank';
        td.appendChild(a);
      } else {
        td.textContent = cell.value;
      }
      tr.appendChild(td);
    });
    table.appendChild(tr);
  });
  // 🛠️ تغليف الجدول بحاوية تسمح بالتمرير
  const tableWrapper = document.createElement('div');
  tableWrapper.style.cssText = 'flex: 1; overflow: auto; width: 100%; height: 100%;';
  tableWrapper.appendChild(table);
  viewer.appendChild(tableWrapper);

  searchInput.addEventListener('input', () => {
    const query = searchInput.value.toLowerCase();
    table.querySelectorAll('tr').forEach(tr => {
      tr.style.display = tr.textContent.toLowerCase().includes(query) ? "" : "none";
    });
  });
}

// إضافة مسار تصفح جوجل درايف
async function loadDriveFolder(folderId) {
  const box = document.getElementById('results');
  box.innerHTML = 'Loading...';
  const r = await fetch(`/api/drive-browse?folder_id=${folderId}`);
  const data = await r.json();
  if (data.error) { box.innerHTML = 'Error: ' + data.error; return; }
  box.innerHTML = '';
  if (currentFolderPath.length > 0) {
    const upBtn = document.createElement('div');
    upBtn.className = 'row';
    upBtn.innerHTML = '⬅ Back';
    upBtn.onclick = () => {
      currentFolderPath.pop();
      const parentId = currentFolderPath.length > 0 ? currentFolderPath[currentFolderPath.length - 1].id : '';
      loadDriveFolder(parentId);
    };
    box.appendChild(upBtn);
  }
  data.folders.forEach(f => {
    const div = document.createElement('div');
    div.className = 'row';
    div.innerHTML = `📁 ${f.name}`;
    div.onclick = () => {
      currentFolderPath.push({ id: f.id, name: f.name });
      loadDriveFolder(f.id);
    };
    box.appendChild(div);
  });
  data.files.forEach(f => {
    const div = document.createElement('div');
    div.className = 'row';
    const icon = (f.ext === '.xlsx' || f.ext === '.xlsm') ? '📊' : '📄';
    div.innerHTML = `${icon} ${f.name}`;
    if (f.ext === '.xlsx' || f.ext === '.xlsm') {
      div.onclick = () => openXlsxTable(f.id);
    } else {
      div.onclick = () => window.open(`/file/${f.id}`, '_blank');
    }
    box.appendChild(div);
  });
}

document.getElementById('q').addEventListener('keydown', e => { if (e.key === 'Enter') doSearch(); });
loadStats();
</script>
</body>
</html>
"""

OVERVIEW_HTML = """
<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>{{ title }}</title>
<style>
  body { background:#020d24; color:#fff; font-family: Tahoma, Arial, sans-serif; margin:0; }
  header { display:flex; align-items:center; justify-content:space-between; padding:10px 20px; background:#03163d; }
  header a { color:#6cf; text-decoration:none; }
  h1 { text-align:center; margin:10px 0; font-size:20px; }
  #grid { columns:5 200px; column-gap:10px; padding:0 15px 10px; max-height: 40vh; overflow-y: auto; }
  .cat { break-inside:avoid; margin-bottom:6px; }
  .box { background:#29b6f6; color:#04203a; font-weight:bold; text-align:center; padding:8px 6px; border-radius:3px; cursor:pointer; font-size:12px; }
  .box:hover { background:#5cd0ff; }
  .box.picked { outline:2px solid #fff; }
  .items { display:flex; flex-wrap:wrap; gap:3px; margin-top:3px; padding:5px; background:#03163d; border-radius:3px; }
  .items button { flex:1 1 60px; background:#0b8fd8; color:#fff; border:none; border-radius:3px; padding:4px 3px; font-size:10px; cursor:pointer; }
  .items button:hover { background:#12a4f5; }
  .items button:disabled { background:#444; opacity:0.5; cursor:not-allowed; }
  #status { text-align:center; color:#9cf; min-height:20px; margin-bottom:5px; }
  #viewer { height: calc(100vh - 220px); margin:0 10px 10px; background:#000; border-radius:8px; display:flex; align-items:center; justify-content:center; overflow:auto; border:1px solid #333; }
  #viewer iframe { width:100%; height:100%; border:none; }
  #viewer img { width:100%; height:100%; object-fit:contain; }
  #fullscreen-btn { cursor:pointer; background:#1f538d; color:#fff; border:none; border-radius:4px; padding:5px 10px; }
</style>
</head>
<body>
<header>
  <a href="/">⬅ Back to Main Dashboard</a>
  <span>🏭 Control and Systems Archive</span>
  <button id="fullscreen-btn" onclick="toggleFullScreen()">⛶ Fullscreen</button>
</header>
<h1>{{ title }}</h1>
<div id="grid"></div>
<div id="status"></div>
<div id="viewer">Click any box to open its map</div>
<script>
async function buildGrid() {
  const grid = document.getElementById('grid');
  grid.textContent = 'Loading...';
  const r = await fetch('/api/overview-data');
  const data = await r.json();
  if (data.error) { grid.textContent = 'Error: ' + data.error; return; }
  grid.textContent = '';
  data.categories.forEach(cat => {
    const catDiv = document.createElement('div');
    catDiv.className = 'cat';
    const box = document.createElement('div');
    box.className = 'box' + (cat.items.length ? '' : ' empty');
    box.textContent = cat.label + (cat.items.length ? ` (${cat.items.length})` : '');
    catDiv.appendChild(box);
    const itemsDiv = document.createElement('div');
    itemsDiv.className = 'items';
    itemsDiv.style.display = 'none';
    cat.items.forEach(item => {
      const b = document.createElement('button');
      b.textContent = item.tag;
      if (item.href) {
        b.onclick = (e) => { e.stopPropagation(); window.open(item.href, '_blank'); };
      } else { b.disabled = true; }
      itemsDiv.appendChild(b);
    });
    catDiv.appendChild(itemsDiv);
    if (cat.items.length) {
      box.onclick = () => {
        const wasOpen = itemsDiv.style.display !== 'none';
        document.querySelectorAll('.items').forEach(d => d.style.display = 'none');
        document.querySelectorAll('.box').forEach(b => b.classList.remove('picked'));
        if (!wasOpen) {
          itemsDiv.style.display = 'flex';
          box.classList.add('picked');
          document.getElementById('status').textContent = `${cat.label} — ${cat.items.length} map(s)`;
        } else { document.getElementById('status').textContent = ''; }
      };
    }
    grid.appendChild(catDiv);
  });
}
function toggleFullScreen() {
  const viewer = document.getElementById('viewer');
  if (!document.fullscreenElement) { viewer.requestFullscreen().catch(err => {}); }
  else { document.exitFullscreen(); }
}
buildGrid();
</script>
</body>
</html>
"""

JUNCTION_HTML = """
<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>{{ title }}</title>
<style>
  body { background:#020d24; color:#fff; font-family: Tahoma, Arial, sans-serif; margin:0; }
  header { display:flex; align-items:center; justify-content:space-between; padding:10px 20px; background:#03163d; }
  header a { color:#6cf; text-decoration:none; }
  h1 { text-align:center; margin:10px 0; font-size:22px; }
  #grid { display: grid; grid-template-columns: repeat(auto-fill, minmax(250px, 1fr)); gap: 15px; padding: 20px; }
  .cat { background: #0f2b5b; border-radius: 12px; padding: 15px; }
  .box { background:#29b6f6; color:#04203a; font-weight:bold; text-align:center; padding:15px; border-radius:8px; margin-bottom:10px; font-size:18px; }
  .items { display: flex; flex-wrap: wrap; gap: 8px; }
  .items button { flex:1 1 100px; background:#0b8fd8; color:#fff; border:none; border-radius:6px; padding:15px; font-size:16px; cursor:pointer; font-weight:bold; }
  .items button:hover { background:#12a4f5; }
  #viewer { display: none; flex: 1; background:#000; flex-direction:column; overflow:hidden; }
  #viewer iframe { width:100%; height:100%; border:none; }
  body.focus-mode #grid, body.focus-mode h1 { display: none !important; }
  body.focus-mode #viewer { display: flex; }
  #exit-focus-btn { position:absolute; top:10px; right:10px; display:none; background:#d9534f; z-index:100; }
  body.focus-mode #exit-focus-btn { display: block; }
  table.xlsx { border-collapse:collapse; font-size:18px; margin:auto; width:100%; }
  table.xlsx td { border:1px solid #444; padding:10px; white-space:nowrap; background:#111; color:#eee; }
</style>
</head>
<body>
<header>
  <a href="/">⬅ Back to Main Dashboard</a>
  <span>🏭 {{ title }}</span>
</header>
<button id="exit-focus-btn" onclick="exitFocusMode()">⬅ Exit Fullscreen</button>
<h1>{{ title }}</h1>
<div id="grid"></div>
<div id="viewer">Click any file to open it</div>

<script>
async function buildGrid() {
  const grid = document.getElementById('grid');
  grid.textContent = 'Loading...';
  const r = await fetch('/api/junction-data');
  const data = await r.json();
  grid.textContent = '';
  data.categories.forEach(cat => {
    const catDiv = document.createElement('div');
    catDiv.className = 'cat';
    const box = document.createElement('div');
    box.className = 'box';
    box.textContent = cat.label + ' (' + cat.items.length + ')';
    catDiv.appendChild(box);
    const itemsDiv = document.createElement('div');
    itemsDiv.className = 'items';
    cat.items.forEach(item => {
      const b = document.createElement('button');
      b.textContent = item.tag;
      if (item.href) {
        b.onclick = (e) => { e.stopPropagation(); enterFocusMode(); openXlsxTable(item.href); };
      } else {
        b.disabled = true;
      }
      itemsDiv.appendChild(b);
    });
    catDiv.appendChild(itemsDiv);
    grid.appendChild(catDiv);
  });
}

function enterFocusMode() { document.body.classList.add('focus-mode'); }
function exitFocusMode() { document.body.classList.remove('focus-mode'); }

async function openXlsxTable(fileId) {
  const viewer = document.getElementById('viewer');
  viewer.innerHTML = 'Loading file from Google Drive...';
  const r = await fetch(`/api/xlsx-table?file_id=${fileId}`);
  const data = await r.json();
  if (data.error) { viewer.textContent = 'Error: ' + data.error; return; }
  viewer.innerHTML = ''; 
  const searchContainer = document.createElement('div');
  searchContainer.style.cssText = 'background: #222; padding: 10px;';
  const searchInput = document.createElement('input');
  searchInput.type = 'text'; searchInput.placeholder = '🔍 Search inside this file...';
  searchInput.style.cssText = 'width: 100%; padding: 10px; background: #333; color: #fff; border: none;';
  searchContainer.appendChild(searchInput);
  viewer.appendChild(searchContainer);
  const table = document.createElement('table');
  table.className = 'xlsx';
  data.rows.forEach(row => {
    const tr = document.createElement('tr');
    row.forEach(cell => {
      const td = document.createElement('td');
      if (cell.style && cell.style.bg) td.style.backgroundColor = cell.style.bg;
      if (cell.href) {
        const a = document.createElement('a');
        a.textContent = cell.value;
        a.href = cell.href;
        a.target = '_blank';
        td.appendChild(a);
      } else {
        td.textContent = cell.value;
      }
      tr.appendChild(td);
    });
    table.appendChild(tr);
  });
  
  // 🛠️ تغليف الجدول بحاوية تسمح بالتمرير
  const tableWrapper = document.createElement('div');
  tableWrapper.style.cssText = 'flex: 1; overflow: auto; width: 100%; height: 100%;';
  tableWrapper.appendChild(table);
  viewer.appendChild(tableWrapper);

  searchInput.addEventListener('input', () => {
    const query = searchInput.value.toLowerCase();
    table.querySelectorAll('tr').forEach(tr => {
      tr.style.display = tr.textContent.toLowerCase().includes(query) ? "" : "none";
    });
  });
}
buildGrid();
</script>
</body>
</html>
"""

# --- Routes الصفحات ---
@app.route("/")
def index():
    return render_template_string(PAGE_HTML)

@app.route("/overview")
def overview():
    return render_template_string(OVERVIEW_HTML, title=OVERVIEW_TITLE)

@app.route("/junction-boxes")
def junction_page():
    return render_template_string(JUNCTION_HTML, title="Junction Boxes")

if __name__ == "__main__":
    threading.Thread(target=background_scan, daemon=True).start()
    app.run(host="0.0.0.0", port=5000, debug=False)
    from flask import Flask, request, jsonify, send_file, render_template_string, abort, session, redirect, url_for
import os
import threading
import time
import openpyxl
from urllib.parse import quote
import io
import json
import google_auth_oauthlib.flow
import drive_helper
import tempfile
import urllib.parse

os.environ['OAUTHLIB_INSECURE_TRANSPORT'] = '1'

app = Flask(__name__, static_folder='static')
app.secret_key = 'super_secret_archive_key'

ROOT_DRIVE_FOLDER_NAME = "DATA SHEET INDEX"
OVERVIEW_TITLE = "R2 PROCESSING AREA"
JUNCTION_FOLDER_NAME = "JUNCTION BOXES"

all_data = []
data_lock = threading.Lock()

# ============ دوال المساعدة ============
def build_hyperlink_map(ws):
    hmap = {}
    try:
        for hl in ws._hyperlinks:
            if hl.ref and hl.target:
                hmap[hl.ref] = hl.target
    except Exception:
        pass
    return hmap

def get_cell_hyperlink_target(cell, hmap):
    if cell.hyperlink and cell.hyperlink.target:
        return cell.hyperlink.target
    target = hmap.get(cell.coordinate)
    if target:
        return target
    return None

# ============ التفويض ============
@app.route('/authorize')
def authorize():
    flow = google_auth_oauthlib.flow.Flow.from_client_secrets_file(
        os.path.join(os.path.dirname(os.path.abspath(__file__)), 'credentials.json'), scopes=drive_helper.SCOPES)
    flow.redirect_uri = url_for('oauth2callback', _external=True)
    authorization_url, state = flow.authorization_url(access_type='offline', include_granted_scopes='true')
    session['state'] = state
    session['code_verifier'] = flow.code_verifier
    return redirect(authorization_url)

@app.route('/oauth2callback')
def oauth2callback():
    state = session['state']
    flow = google_auth_oauthlib.flow.Flow.from_client_secrets_file(
        os.path.join(os.path.dirname(os.path.abspath(__file__)), 'credentials.json'), scopes=drive_helper.SCOPES, state=state)
    flow.redirect_uri = url_for('oauth2callback', _external=True)
    authorization_response = request.url
    code_verifier = session.get('code_verifier')
    flow.fetch_token(authorization_response=authorization_response, code_verifier=code_verifier)
    creds = flow.credentials
    with open(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'token.json'), 'w') as token:
        token.write(creds.to_json())
    return redirect(url_for('index'))

# ============ المسح في الخلفية ============
def scan_drive_recursive(service, folder_id, current_path=""):
    files = drive_helper.list_files_in_folder(service, folder_id)
    for f in files:
        if f['mimeType'] == 'application/vnd.google-apps.folder':
            scan_drive_recursive(service, f['id'], current_path + f['name'] + "/")
        elif f['name'].endswith(('.xlsx', '.xlsm')) and not f['name'].startswith('~$'):
            try:
                file_bytes = drive_helper.download_file(service, f['id'])
                wb = openpyxl.load_workbook(file_bytes, data_only=True)
                ws = wb.active
                hmap = build_hyperlink_map(ws)
                
                for row in ws.iter_rows():
                    for cell in row:
                        target = get_cell_hyperlink_target(cell, hmap)
                        if cell.value and target:
                            with data_lock:
                                all_data.append({
                                    "tag": str(cell.value).strip(),
                                    "path": current_path + f['name'],
                                    "file_id": f['id'],
                                    "folder": current_path.split('/')[0] if current_path else "Root"
                                })
            except Exception as e:
                print(f"Error reading {f['name']}: {e}")

def background_scan():
    while True:
        service = drive_helper.get_drive_service()
        if service:
            print("Scanning Google Drive... Please wait.")
            root_id = drive_helper.find_folder_id(service, ROOT_DRIVE_FOLDER_NAME)
            if root_id:
                with data_lock:
                    all_data.clear()
                scan_drive_recursive(service, root_id)
                print(f"Scan complete. Found {len(all_data)} tags.")
            else:
                print(f"Root folder '{ROOT_DRIVE_FOLDER_NAME}' not found in Google Drive.")
        else:
            print("Waiting for authorization...")
        time.sleep(3600)

# ============ API Routes ============
@app.route("/api/stats")
def api_stats():
    with data_lock:
        data = list(all_data)
    folders = {}
    for d in data:
        folders[d["folder"]] = folders.get(d["folder"], 0) + 1
    return jsonify({
        "total": len(data),
        "folder_count": len(folders),
        "excel_files": len(set(d["path"] for d in data)),
        "folders": folders
    })

@app.route("/api/search")
def api_search():
    q = request.args.get("q", "").lower()
    with data_lock:
        data = list(all_data)
    results = []
    for i, d in enumerate(data):
        if q in d["tag"].lower():
            results.append({"id": i, "tag": d["tag"], "sheet_no": "", "file_id": d["file_id"]})
    return jsonify(results)

@app.route("/file/<file_id>")
def serve_drive_file(file_id):
    service = drive_helper.get_drive_service()
    if not service: abort(401)
    try:
        file_meta = service.files().get(fileId=file_id, fields="name, mimeType").execute()
        file_bytes = drive_helper.download_file(service, file_id)
        temp_path = os.path.join(tempfile.gettempdir(), file_meta['name'])
        with open(temp_path, 'wb') as f:
            f.write(file_bytes.getvalue())
        return send_file(temp_path, mimetype=file_meta['mimeType'])
    except Exception as e:
        return str(e), 500

@app.route("/api/xlsx-table")
def api_xlsx_table():
    file_id = request.args.get("file_id")
    if not file_id: abort(404)
    
    service = drive_helper.get_drive_service()
    if not service: return jsonify({"error": "Not authorized"}), 401
    
    try:
        file_meta = service.files().get(fileId=file_id, fields="name, parents").execute()
        excel_parent_id = file_meta['parents'][0] if 'parents' in file_meta else None
        file_bytes = drive_helper.download_file(service, file_id)
        wb = openpyxl.load_workbook(file_bytes)
        ws = wb.active
        hmap = build_hyperlink_map(ws)
    except Exception as e:
        print(f"❌ ERROR downloading/reading file {file_id}: {e}")
        return jsonify({"error": str(e)}), 500

    raw_rows = list(ws.iter_rows())
    clean_grid = []
    
    for row in raw_rows:
        clean_row = []
        for cell in row:
            href = None
            target = get_cell_hyperlink_target(cell, hmap)
            if target:
                decoded_target = urllib.parse.unquote(target.replace("\\", "/"))
                if '#' in decoded_target:
                    decoded_target = decoded_target.split('#')[0]
                href = f"/find-and-open?target={quote(decoded_target)}&parent_id={excel_parent_id}"
                    
            val = "" if cell.value is None else str(cell.value).strip()
            if val or href:
                style = {}
                try:
                    if cell.fill.patternType == "solid":
                        fill_color = cell.fill.start_color
                        if fill_color.rgb and fill_color.rgb != "00000000": style["bg"] = f"#{fill_color.rgb[2:]}"
                except: pass
                clean_row.append({"value": val, "href": href, "style": style})
        if clean_row: clean_grid.append(clean_row)
        
    return jsonify({"rows": clean_grid})

@app.route("/find-and-open")
def find_and_open():
    target = request.args.get("target")
    parent_id = request.args.get("parent_id")
    if not target: abort(400)
    
    service = drive_helper.get_drive_service()
    if not service: abort(401)
    
    path_parts = [p for p in target.split('/') if p and p != '..']
    filename = path_parts[-1]
    linked_file_id = None
    
    if parent_id and len(path_parts) > 1:
        linked_file_id = drive_helper.find_file_by_path(service, path_parts, parent_id)
    if not linked_file_id and parent_id:
        linked_file_id = drive_helper.find_file_in_tree(service, filename, parent_id)
    if not linked_file_id:
        linked_file_id = drive_helper.find_file_by_name(service, filename)
        
    if linked_file_id:
        return redirect(f"/file/{linked_file_id}")
    else:
        return f"<h1>File Not Found</h1><p>The file '{filename}' was not found.</p>", 404

@app.route("/api/junction-data")
def api_junction_data():
    service = drive_helper.get_drive_service()
    if not service: return jsonify({"error": "Unauthorized"}), 401
    root_id = drive_helper.find_folder_id(service, ROOT_DRIVE_FOLDER_NAME)
    if not root_id: return jsonify({"categories": []})
    folder_id = drive_helper.find_folder_id(service, JUNCTION_FOLDER_NAME, parent_id=root_id)
    if not folder_id: return jsonify({"categories": []})
    files = drive_helper.list_files_in_folder(service, folder_id)
    items = []
    for f in files:
        if f['name'].lower().endswith(('.xlsx', '.xlsm')):
            items.append({"tag": f['name'].split('.')[0], "href": f['id']})
    categories = [{"label": "Junction Boxes Files", "items": items}] if items else []
    return jsonify({"categories": categories})

@app.route("/api/overview-data")
def api_overview_data():
    service = drive_helper.get_drive_service()
    if not service: return jsonify({"error": "Unauthorized"}), 401
    root_id = drive_helper.find_folder_id(service, ROOT_DRIVE_FOLDER_NAME)
    if not root_id: return jsonify({"error": "Root not found"}), 404
    pi_id = drive_helper.find_folder_id(service, "P&I", parent_id=root_id)
    if not pi_id: return jsonify({"error": "P&I folder not found"}), 404
    files = drive_helper.list_files_in_folder(service, pi_id)
    pi_file_id = next((f['id'] for f in files if f['name'].lower() == 'p&i.xlsx'), None)
    if not pi_file_id: return jsonify({"error": "P&I.xlsx not found"}), 404
    
    try:
        file_bytes = drive_helper.download_file(service, pi_file_id)
        wb = openpyxl.load_workbook(file_bytes, data_only=True)
        ws = wb.active
        hmap = build_hyperlink_map(ws)
        categories = []
        for row in ws.iter_rows():
            current_label = None
            current_items = []
            def flush():
                if current_label: categories.append({"label": current_label, "items": current_items})
            for cell in row:
                target = get_cell_hyperlink_target(cell, hmap)
                if target:
                    if current_label:
                        filename = os.path.basename(target.replace("\\", "/"))
                        linked_id = drive_helper.find_file_by_name(service, filename)
                        href = f"/file/{linked_id}" if linked_id else None
                        current_items.append({"tag": str(cell.value).strip() if cell.value else "?", "href": href})
                elif cell.value not in (None, ""):
                    flush(); current_label = str(cell.value).strip(); current_items = []
                else:
                    flush(); current_label = None; current_items = []
            flush()
        return jsonify({"categories": categories})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

# ============ القوالب (HTML) ============
@app.route("/api/drive-browse")
def api_drive_browse():
    service = drive_helper.get_drive_service()
    if not service: return jsonify({"error": "Unauthorized"}), 401

    folder_id = request.args.get("folder_id")
    
    try:
        if not folder_id:
            folder_id = drive_helper.find_folder_id(service, ROOT_DRIVE_FOLDER_NAME)
            if not folder_id: return jsonify({"error": "Root folder not found"}), 404

        files = drive_helper.list_files_in_folder(service, folder_id)
        folders_out = []
        files_out = []

        for f in files:
            if f['mimeType'] == 'application/vnd.google-apps.folder':
                folders_out.append({"id": f['id'], "name": f['name']})
            else:
                ext = os.path.splitext(f['name'])[1].lower()
                if ext in ('.xlsx', '.xlsm', '.pdf', '.jpg', '.jpeg', '.png', '.drawio'):
                    files_out.append({"id": f['id'], "name": f['name'], "ext": ext})

        folders_out.sort(key=lambda x: x['name'])
        files_out.sort(key=lambda x: x['name'])

        return jsonify({"folders": folders_out, "files": files_out})
    except Exception as e:
        print(f"❌ Browse Error: {e}")
        return jsonify({"error": str(e)}), 500
PAGE_HTML = """
<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>Control and Systems Archive</title>
<script src="/static/echarts.min.js"></script>
<script src="/static/echarts-gl.min.js"></script>
<style>
  html, body { height:100%; margin:0; }
  body { background:#1a1a1a; color:#eee; font-family: Tahoma, Arial, sans-serif; display:flex; height:100vh; }
  #sidebar { width:300px; background:#232323; padding:16px; box-sizing:border-box; display:flex; flex-direction:column; }
  #sidebar h1 { font-size:16px; margin:0 0 16px; }
  input[type=text] { width:100%; padding:8px; border-radius:6px; border:none; background:#2b2b2b; color:#fff; box-sizing:border-box; }
  button { cursor:pointer; border:none; border-radius:6px; background:#1f538d; color:#fff; padding:8px 12px; margin-top:8px; }
  button:hover { background:#2a6bb0; }
  .tabs { display:flex; gap:6px; margin-top:12px; }
  .tab { flex:1; background:#333; text-align:center; }
  .tab.active { background:#1f538d; }
  #results { flex:1; overflow-y:auto; margin-top:10px; background:#2b2b2b; border-radius:6px; }
  .row { padding:8px 10px; border-bottom:1px solid #333; cursor:pointer; font-size:13px; }
  .row:hover { background:#1f538d; }
  #main { flex:1; display:flex; flex-direction:column; padding:16px; box-sizing:border-box; overflow:hidden; }
  .cards { display:flex; gap:10px; margin-bottom:16px; flex-shrink:0; }
  .card { flex:1; border-radius:10px; padding:20px; text-align:center; color:#fff; font-weight:bold; font-size:20px; }
  .card b { display:block; font-size:42px; margin-top:8px; }
  .card:nth-child(1) { background:#00bcd4; }
  .card:nth-child(2) { background:#2ecc71; }
  .card:nth-child(3) { background:#ffa726; }
  .chart-container { flex-grow:1; background:#2b2b2b; border-radius:8px; overflow:hidden; position:relative; }
  #folder-chart { width:100%; height:100%; }
  #viewer { display: none; flex:1; background:#000; border-radius:8px; flex-direction:column; overflow:hidden; }
  #viewer iframe { width:100%; height:100%; border:none; }
  body.focus-mode .chart-container, body.focus-mode .cards { display: none !important; }
  body.focus-mode #viewer { display: flex; }
  #exit-focus-btn { position:absolute; top:20px; right:20px; display:none; background:#d9534f; z-index:100; }
  body.focus-mode #exit-focus-btn { display: block; }
  table.xlsx { border-collapse:collapse; font-size:18px; margin:auto; width:100%; }
  table.xlsx td { border:1px solid #444; padding:10px; white-space:nowrap; background:#111; color:#eee; }
  table.xlsx a { color:#6cf; text-decoration:underline; cursor:pointer; }
</style>
</head>
<body>
<div id="sidebar">
  <h1>🏭 Archive (Google Drive)</h1>
  <div class="tabs">
    <button class="tab active" id="tab-search" onclick="switchTab('search')">🔍 Search</button>
    <button class="tab" id="tab-browse" onclick="switchTab('browse')">🗂 Browse</button>
  </div>
  <button style="width:100%;" onclick="location.href='/overview'">🗺️ R2 General Map</button>
  <button style="width:100%;" onclick="location.href='/junction-boxes'">🔌 Junction Boxes</button>
  <div id="search-box">
    <input type="text" id="q" placeholder="Search Tag...">
    <button onclick="doSearch()">🔍 Search</button>
  </div>
  <div id="results"></div>
</div>
<div id="main">
  <button id="exit-focus-btn" onclick="exitFocusMode()">⬅ Exit Fullscreen</button>
  <div class="cards">
    <div class="card">Total Tags<b id="c-total">0</b></div>
    <div class="card">Total Folders<b id="c-folders">0</b></div>
    <div class="card">Excel Files<b id="c-excel">0</b></div>
  </div>
  <div class="chart-container"><div id="folder-chart"></div></div>
  <div id="viewer">Select a file to open.</div>
</div>

<script>
let mode = 'search';
let currentFolderPath = [];

function switchTab(which) {
  mode = which;
  document.getElementById('tab-search').classList.toggle('active', which === 'search');
  document.getElementById('tab-browse').classList.toggle('active', which === 'browse');
  document.getElementById('search-box').style.display = which === 'search' ? 'block' : 'none';
  document.getElementById('results').innerHTML = '';
  if (which === 'search') { doSearch(); }
  else { currentFolderPath = []; loadDriveFolder(''); }
}

async function loadDriveFolder(folderId) {
  const box = document.getElementById('results');
  box.innerHTML = 'Loading...';
  const r = await fetch(`/api/drive-browse?folder_id=${folderId}`);
  const data = await r.json();
  if (data.error) { box.innerHTML = 'Error: ' + data.error; return; }
  box.innerHTML = '';
  if (currentFolderPath.length > 0) {
    const upBtn = document.createElement('div');
    upBtn.className = 'row';
    upBtn.innerHTML = '⬅ Back';
    upBtn.onclick = () => {
      currentFolderPath.pop();
      const parentId = currentFolderPath.length > 0 ? currentFolderPath[currentFolderPath.length - 1].id : '';
      loadDriveFolder(parentId);
    };
    box.appendChild(upBtn);
  }
  data.folders.forEach(f => {
    const div = document.createElement('div');
    div.className = 'row';
    div.innerHTML = `📁 ${f.name}`;
    div.onclick = () => {
      currentFolderPath.push({ id: f.id, name: f.name });
      loadDriveFolder(f.id);
    };
    box.appendChild(div);
  });
  data.files.forEach(f => {
    const div = document.createElement('div');
    div.className = 'row';
    const icon = (f.ext === '.xlsx' || f.ext === '.xlsm') ? '📊' : '📄';
    div.innerHTML = `${icon} ${f.name}`;
    if (f.ext === '.xlsx' || f.ext === '.xlsm') {
      div.onclick = () => openXlsxTable(f.id);
    } else {
      div.onclick = () => window.open(`/file/${f.id}`, '_blank');
    }
    box.appendChild(div);
  });
}

async function loadStats() {
  const r = await fetch('/api/stats');
  const s = await r.json();
  document.getElementById('c-total').textContent = s.total;
  document.getElementById('c-folders').textContent = s.folder_count;
  document.getElementById('c-excel').textContent = s.excel_files;
  const chartDom = document.getElementById('folder-chart');
  const myChart = echarts.init(chartDom);
  const folders = Object.keys(s.folders).map(name => [name, s.folders[name], s.folders[name]]);
  folders.sort((a, b) => b[1] - a[1]);
  myChart.setOption({
    tooltip: {}, visualMap: { max: Math.max(...Object.values(s.folders)), inRange: { color: ['#00bcd4', '#ffa726'] } },
    xAxis3D: { type: 'category', data: folders.map(f => f[0]), axisLabel: { rotate: 45, color: '#fff', fontSize: 10 } },
    yAxis3D: { type: 'value' }, zAxis3D: { type: 'value' },
    grid3D: { boxWidth: 200, boxDepth: 80, viewControl: { autoRotate: true } },
    series: [{ type: 'bar3D', data: folders, shading: 'lambert', barSize: 10 }]
  });
  window.addEventListener('resize', () => myChart.resize());
}

function addRow(text, onClick) {
  const div = document.createElement('div'); div.className = 'row'; div.textContent = text; div.onclick = onClick;
  document.getElementById('results').appendChild(div);
}

async function doSearch() {
  const q = document.getElementById('q').value.trim();
  const box = document.getElementById('results');
  if (!q) { box.innerHTML = ''; return; }
  const r = await fetch(`/api/search?q=${encodeURIComponent(q)}`);
  const files = await r.json();
  box.innerHTML = '';
  files.forEach(f => addRow(f.tag, () => openXlsxTable(f.file_id)));
}

function enterFocusMode() { document.body.classList.add('focus-mode'); }
function exitFocusMode() { document.body.classList.remove('focus-mode'); }

async function openXlsxTable(fileId) {
  enterFocusMode();
  const viewer = document.getElementById('viewer');
  viewer.innerHTML = 'Loading file from Google Drive...';
  const r = await fetch(`/api/xlsx-table?file_id=${fileId}`);
  const data = await r.json();
  if (data.error) { viewer.textContent = 'Error: ' + data.error; return; }
  viewer.innerHTML = '';
  const searchContainer = document.createElement('div');
  searchContainer.style.cssText = 'background: #222; padding: 10px;';
  const searchInput = document.createElement('input');
  searchInput.type = 'text'; searchInput.placeholder = '🔍 Search inside this file...';
  searchInput.style.cssText = 'width: 100%; padding: 10px; background: #333; color: #fff; border: none;';
  searchContainer.appendChild(searchInput);
  viewer.appendChild(searchContainer);
  const table = document.createElement('table');
  table.className = 'xlsx';
  data.rows.forEach(row => {
    const tr = document.createElement('tr');
    row.forEach(cell => {
      const td = document.createElement('td');
      if (cell.style && cell.style.bg) td.style.backgroundColor = cell.style.bg;
      if (cell.href) {
        const a = document.createElement('a');
        a.textContent = cell.value;
        a.href = cell.href;
        a.target = '_blank';
        td.appendChild(a);
      } else {
        td.textContent = cell.value;
      }
      tr.appendChild(td);
    });
    table.appendChild(tr);
  });
  // 🛠️ تغليف الجدول بحاوية تسمح بالتمرير
  const tableWrapper = document.createElement('div');
  tableWrapper.style.cssText = 'flex: 1; overflow: auto; width: 100%; height: 100%;';
  tableWrapper.appendChild(table);
  viewer.appendChild(tableWrapper);

  searchInput.addEventListener('input', () => {
    const query = searchInput.value.toLowerCase();
    table.querySelectorAll('tr').forEach(tr => {
      tr.style.display = tr.textContent.toLowerCase().includes(query) ? "" : "none";
    });
  });
}

// إضافة مسار تصفح جوجل درايف
async function loadDriveFolder(folderId) {
  const box = document.getElementById('results');
  box.innerHTML = 'Loading...';
  const r = await fetch(`/api/drive-browse?folder_id=${folderId}`);
  const data = await r.json();
  if (data.error) { box.innerHTML = 'Error: ' + data.error; return; }
  box.innerHTML = '';
  if (currentFolderPath.length > 0) {
    const upBtn = document.createElement('div');
    upBtn.className = 'row';
    upBtn.innerHTML = '⬅ Back';
    upBtn.onclick = () => {
      currentFolderPath.pop();
      const parentId = currentFolderPath.length > 0 ? currentFolderPath[currentFolderPath.length - 1].id : '';
      loadDriveFolder(parentId);
    };
    box.appendChild(upBtn);
  }
  data.folders.forEach(f => {
    const div = document.createElement('div');
    div.className = 'row';
    div.innerHTML = `📁 ${f.name}`;
    div.onclick = () => {
      currentFolderPath.push({ id: f.id, name: f.name });
      loadDriveFolder(f.id);
    };
    box.appendChild(div);
  });
  data.files.forEach(f => {
    const div = document.createElement('div');
    div.className = 'row';
    const icon = (f.ext === '.xlsx' || f.ext === '.xlsm') ? '📊' : '📄';
    div.innerHTML = `${icon} ${f.name}`;
    if (f.ext === '.xlsx' || f.ext === '.xlsm') {
      div.onclick = () => openXlsxTable(f.id);
    } else {
      div.onclick = () => window.open(`/file/${f.id}`, '_blank');
    }
    box.appendChild(div);
  });
}

document.getElementById('q').addEventListener('keydown', e => { if (e.key === 'Enter') doSearch(); });
loadStats();
</script>
</body>
</html>
"""

OVERVIEW_HTML = """
<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>{{ title }}</title>
<style>
  body { background:#020d24; color:#fff; font-family: Tahoma, Arial, sans-serif; margin:0; }
  header { display:flex; align-items:center; justify-content:space-between; padding:10px 20px; background:#03163d; }
  header a { color:#6cf; text-decoration:none; }
  h1 { text-align:center; margin:10px 0; font-size:20px; }
  #grid { columns:5 200px; column-gap:10px; padding:0 15px 10px; max-height: 40vh; overflow-y: auto; }
  .cat { break-inside:avoid; margin-bottom:6px; }
  .box { background:#29b6f6; color:#04203a; font-weight:bold; text-align:center; padding:8px 6px; border-radius:3px; cursor:pointer; font-size:12px; }
  .box:hover { background:#5cd0ff; }
  .box.picked { outline:2px solid #fff; }
  .items { display:flex; flex-wrap:wrap; gap:3px; margin-top:3px; padding:5px; background:#03163d; border-radius:3px; }
  .items button { flex:1 1 60px; background:#0b8fd8; color:#fff; border:none; border-radius:3px; padding:4px 3px; font-size:10px; cursor:pointer; }
  .items button:hover { background:#12a4f5; }
  .items button:disabled { background:#444; opacity:0.5; cursor:not-allowed; }
  #status { text-align:center; color:#9cf; min-height:20px; margin-bottom:5px; }
  #viewer { height: calc(100vh - 220px); margin:0 10px 10px; background:#000; border-radius:8px; display:flex; align-items:center; justify-content:center; overflow:auto; border:1px solid #333; }
  #viewer iframe { width:100%; height:100%; border:none; }
  #viewer img { width:100%; height:100%; object-fit:contain; }
  #fullscreen-btn { cursor:pointer; background:#1f538d; color:#fff; border:none; border-radius:4px; padding:5px 10px; }
</style>
</head>
<body>
<header>
  <a href="/">⬅ Back to Main Dashboard</a>
  <span>🏭 Control and Systems Archive</span>
  <button id="fullscreen-btn" onclick="toggleFullScreen()">⛶ Fullscreen</button>
</header>
<h1>{{ title }}</h1>
<div id="grid"></div>
<div id="status"></div>
<div id="viewer">Click any box to open its map</div>
<script>
async function buildGrid() {
  const grid = document.getElementById('grid');
  grid.textContent = 'Loading...';
  const r = await fetch('/api/overview-data');
  const data = await r.json();
  if (data.error) { grid.textContent = 'Error: ' + data.error; return; }
  grid.textContent = '';
  data.categories.forEach(cat => {
    const catDiv = document.createElement('div');
    catDiv.className = 'cat';
    const box = document.createElement('div');
    box.className = 'box' + (cat.items.length ? '' : ' empty');
    box.textContent = cat.label + (cat.items.length ? ` (${cat.items.length})` : '');
    catDiv.appendChild(box);
    const itemsDiv = document.createElement('div');
    itemsDiv.className = 'items';
    itemsDiv.style.display = 'none';
    cat.items.forEach(item => {
      const b = document.createElement('button');
      b.textContent = item.tag;
      if (item.href) {
        b.onclick = (e) => { e.stopPropagation(); window.open(item.href, '_blank'); };
      } else { b.disabled = true; }
      itemsDiv.appendChild(b);
    });
    catDiv.appendChild(itemsDiv);
    if (cat.items.length) {
      box.onclick = () => {
        const wasOpen = itemsDiv.style.display !== 'none';
        document.querySelectorAll('.items').forEach(d => d.style.display = 'none');
        document.querySelectorAll('.box').forEach(b => b.classList.remove('picked'));
        if (!wasOpen) {
          itemsDiv.style.display = 'flex';
          box.classList.add('picked');
          document.getElementById('status').textContent = `${cat.label} — ${cat.items.length} map(s)`;
        } else { document.getElementById('status').textContent = ''; }
      };
    }
    grid.appendChild(catDiv);
  });
}
function toggleFullScreen() {
  const viewer = document.getElementById('viewer');
  if (!document.fullscreenElement) { viewer.requestFullscreen().catch(err => {}); }
  else { document.exitFullscreen(); }
}
buildGrid();
</script>
</body>
</html>
"""

JUNCTION_HTML = """
<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>{{ title }}</title>
<style>
  body { background:#020d24; color:#fff; font-family: Tahoma, Arial, sans-serif; margin:0; }
  header { display:flex; align-items:center; justify-content:space-between; padding:10px 20px; background:#03163d; }
  header a { color:#6cf; text-decoration:none; }
  h1 { text-align:center; margin:10px 0; font-size:22px; }
  #grid { display: grid; grid-template-columns: repeat(auto-fill, minmax(250px, 1fr)); gap: 15px; padding: 20px; }
  .cat { background: #0f2b5b; border-radius: 12px; padding: 15px; }
  .box { background:#29b6f6; color:#04203a; font-weight:bold; text-align:center; padding:15px; border-radius:8px; margin-bottom:10px; font-size:18px; }
  .items { display: flex; flex-wrap: wrap; gap: 8px; }
  .items button { flex:1 1 100px; background:#0b8fd8; color:#fff; border:none; border-radius:6px; padding:15px; font-size:16px; cursor:pointer; font-weight:bold; }
  .items button:hover { background:#12a4f5; }
  #viewer { display: none; flex: 1; background:#000; flex-direction:column; overflow:hidden; }
  #viewer iframe { width:100%; height:100%; border:none; }
  body.focus-mode #grid, body.focus-mode h1 { display: none !important; }
  body.focus-mode #viewer { display: flex; }
  #exit-focus-btn { position:absolute; top:10px; right:10px; display:none; background:#d9534f; z-index:100; }
  body.focus-mode #exit-focus-btn { display: block; }
  table.xlsx { border-collapse:collapse; font-size:18px; margin:auto; width:100%; }
  table.xlsx td { border:1px solid #444; padding:10px; white-space:nowrap; background:#111; color:#eee; }
</style>
</head>
<body>
<header>
  <a href="/">⬅ Back to Main Dashboard</a>
  <span>🏭 {{ title }}</span>
</header>
<button id="exit-focus-btn" onclick="exitFocusMode()">⬅ Exit Fullscreen</button>
<h1>{{ title }}</h1>
<div id="grid"></div>
<div id="viewer">Click any file to open it</div>

<script>
async function buildGrid() {
  const grid = document.getElementById('grid');
  grid.textContent = 'Loading...';
  const r = await fetch('/api/junction-data');
  const data = await r.json();
  grid.textContent = '';
  data.categories.forEach(cat => {
    const catDiv = document.createElement('div');
    catDiv.className = 'cat';
    const box = document.createElement('div');
    box.className = 'box';
    box.textContent = cat.label + ' (' + cat.items.length + ')';
    catDiv.appendChild(box);
    const itemsDiv = document.createElement('div');
    itemsDiv.className = 'items';
    cat.items.forEach(item => {
      const b = document.createElement('button');
      b.textContent = item.tag;
      if (item.href) {
        b.onclick = (e) => { e.stopPropagation(); enterFocusMode(); openXlsxTable(item.href); };
      } else {
        b.disabled = true;
      }
      itemsDiv.appendChild(b);
    });
    catDiv.appendChild(itemsDiv);
    grid.appendChild(catDiv);
  });
}

function enterFocusMode() { document.body.classList.add('focus-mode'); }
function exitFocusMode() { document.body.classList.remove('focus-mode'); }

async function openXlsxTable(fileId) {
  const viewer = document.getElementById('viewer');
  viewer.innerHTML = 'Loading file from Google Drive...';
  const r = await fetch(`/api/xlsx-table?file_id=${fileId}`);
  const data = await r.json();
  if (data.error) { viewer.textContent = 'Error: ' + data.error; return; }
  viewer.innerHTML = ''; 
  const searchContainer = document.createElement('div');
  searchContainer.style.cssText = 'background: #222; padding: 10px;';
  const searchInput = document.createElement('input');
  searchInput.type = 'text'; searchInput.placeholder = '🔍 Search inside this file...';
  searchInput.style.cssText = 'width: 100%; padding: 10px; background: #333; color: #fff; border: none;';
  searchContainer.appendChild(searchInput);
  viewer.appendChild(searchContainer);
  const table = document.createElement('table');
  table.className = 'xlsx';
  data.rows.forEach(row => {
    const tr = document.createElement('tr');
    row.forEach(cell => {
      const td = document.createElement('td');
      if (cell.style && cell.style.bg) td.style.backgroundColor = cell.style.bg;
      if (cell.href) {
        const a = document.createElement('a');
        a.textContent = cell.value;
        a.href = cell.href;
        a.target = '_blank';
        td.appendChild(a);
      } else {
        td.textContent = cell.value;
      }
      tr.appendChild(td);
    });
    table.appendChild(tr);
  });
  
  // 🛠️ تغليف الجدول بحاوية تسمح بالتمرير
  const tableWrapper = document.createElement('div');
  tableWrapper.style.cssText = 'flex: 1; overflow: auto; width: 100%; height: 100%;';
  tableWrapper.appendChild(table);
  viewer.appendChild(tableWrapper);

  searchInput.addEventListener('input', () => {
    const query = searchInput.value.toLowerCase();
    table.querySelectorAll('tr').forEach(tr => {
      tr.style.display = tr.textContent.toLowerCase().includes(query) ? "" : "none";
    });
  });
}
buildGrid();
</script>
</body>
</html>
"""

# --- Routes الصفحات ---
@app.route("/")
def index():
    return render_template_string(PAGE_HTML)

@app.route("/overview")
def overview():
    return render_template_string(OVERVIEW_HTML, title=OVERVIEW_TITLE)

@app.route("/junction-boxes")
def junction_page():
    return render_template_string(JUNCTION_HTML, title="Junction Boxes")

if __name__ == "__main__":
    threading.Thread(target=background_scan, daemon=True).start()
    app.run(host="0.0.0.0", port=5000, debug=False)
